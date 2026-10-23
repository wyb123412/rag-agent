from pathlib import Path

from rag.load import load_documents


def test_load_md_and_txt(tmp_path):
    (tmp_path / "a.md").write_text("# 标题\n内容", encoding="utf-8")
    (tmp_path / "b.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "c.pdf").write_bytes(b"not a real pdf")  # 无效 pdf 应被跳过
    docs = load_documents(str(tmp_path))
    names = sorted(Path(d["path"]).name for d in docs)
    assert names == ["a.md", "b.txt"]


def test_empty_file_skipped(tmp_path):
    (tmp_path / "empty.md").write_text("   \n", encoding="utf-8")
    assert load_documents(str(tmp_path)) == []


def test_load_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    doc.save(str(tmp_path / "笔记.docx"))
    docs = load_documents(str(tmp_path))
    assert len(docs) == 1
    assert "第一段内容" in docs[0]["text"]
    assert "第二段内容" in docs[0]["text"]


def test_unsupported_suffix_skipped_and_recursion(tmp_path):
    (tmp_path / "photo.jpg").write_bytes(b"\xff\xd8\xff\xe0")  # 不支持的后缀应跳过
    sub = tmp_path / "子目录"
    sub.mkdir()
    (sub / "内层.md").write_text("嵌套内容", encoding="utf-8")
    docs = load_documents(str(tmp_path))
    names = sorted(Path(d["path"]).name for d in docs)
    assert names == ["内层.md"]  # 递归到子目录，且跳过 .jpg
